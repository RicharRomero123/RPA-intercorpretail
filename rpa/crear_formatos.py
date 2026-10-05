"""Crea las plantillas en Excel que se descargan desde el botón «Cargar …» de la web (web/public/formatos).

Cada plantilla trae la hoja de datos con los títulos exactos que lee el cargador (web/src/lib/cargas.ts, mismas columnas que
web/src/lib/formatos.ts) y una hoja «Instrucciones» con qué poner en cada columna y un ejemplo. La hoja de datos va vacía a
propósito: así no se carga por error una fila de ejemplo.

Uso:  python crear_formatos.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

DESTINO = Path(__file__).resolve().parent.parent / "web" / "public" / "formatos"
FILAS = 2000  # filas con validación y formato preparados

TITULO = Font(name="Arial", bold=True, color="FFFFFF", size=10)
FONDO = PatternFill("solid", fgColor="7A4A22")
FONDO_OPC = PatternFill("solid", fgColor="A8805C")
NORMAL = Font(name="Arial", size=10)
NEGRITA = Font(name="Arial", size=10, bold=True)
LINEA = Border(bottom=Side(style="thin", color="BDBDBD"))

# Cada columna: (título, obligatoria, qué poner, ejemplo, formato de número, ancho, lista de valores permitidos)
TIENDAS = [
    ("SUCURSAL", True, "Nombre de la tienda", "Surco", None, 14,
     ["Abancay", "Jose Granda", "Pana Norte", "Sol de Oro", "Surco", "Trinidad", "Tupac"]),
    ("CANAL", True, "Canal de venta", "Tienda", None, 10, None),
    ("Fecha", True, "Día de la venta (dd/mm/aaaa)", "20/09/2026", "dd/mm/yyyy", 12, None),
    ("CODIGO", True, "Código del producto en el sistema de la tienda", "TUR1114", None, 12, None),
    ("COD. SECUNDARIO", True, "Código corto del producto", "TK", None, 16, None),
    ("Tipo de Precio", True, "Tipo de precio", "Unidad", None, 14, None),
    ("Categoría-Cliente", True, "Categoría del cliente", "Estándar", None, 16, None),
    ("Tienda-Cantidad", True, "Unidades vendidas ese día", "35", "#,##0", 15, None),
    ("Tienda-Total Venta S/", True, "Venta en soles, con IGV", "542.50", "#,##0.00", 19, None),
]
RETAIL = [
    ("RAZON SOCIAL EMISOR", False, "Empresa que factura", "SFC alimentos", None, 18, None),
    ("RUC", False, "RUC del cliente", "20100070970", "@", 13, None),
    ("RAZON SOCIAL CLIENTE", True, "Cadena a la que se despacha", "TOTTUS", None, 24, None),
    ("CANTIDAD", True, "Unidades despachadas", "3528", "#,##0", 11, None),
    ("PRECIO UNITARIO", False, "Precio por unidad en soles", "11.38", "#,##0.00", 15, None),
    ("TIPO DE VENTA", False, "Tipo de venta", "RETAIL", None, 13, None),
    ("SKU", True, "Código oficial del producto", "TMA1113", None, 11, None),
    ("PRODUCTO", False, "Nombre del producto", "Turrón 450g ajonjolí", None, 22, None),
    ("MONTO CANCELADO", True, "Monto total del despacho en soles", "40148.64", "#,##0.00", 16, None),
    ("CONDICION DE PAGO", False, "Crédito o contado", "Credito", None, 17, ["Credito", "Contado"]),
    ("DIRECCION", False, "Dirección de entrega", "CD Secos Huachipa", None, 24, None),
    ("DETALLE DE DESPACHO", False, "Qué se envió", "Caja master + rótulo", None, 22, None),
    ("DIA DE DESPACHO", True, "Día del despacho (dd/mm/aaaa)", "23/09/2026", "dd/mm/yyyy", 15, None),
    ("STATUS", False, "Estado del despacho", "Entregado", None, 12, ["Entregado", "Pendiente", "Anulado"]),
]
VIRTUAL = [
    ("MES", False, "Mes (solo referencia)", "Septiembre", None, 11, None),
    ("Fecha Registro", True, "Día del pago (dd/mm/aaaa)", "24/09/2026", "dd/mm/yyyy", 14, None),
    ("Nro Comprobante", True, "Boleta o factura de ContaNet, con serie y número", "B008-4837", None, 16, None),
    ("Tercero", False, "Nombre del cliente", "Milagros Olortegui", None, 22, None),
    ("Documento", False, "DNI o RUC del cliente", "45879632", "@", 12, None),
    ("Código", False, "Código del producto", "TT1114", None, 10, None),
    ("Descripción", False, "Producto, siempre escrito igual (manda sobre el código)", "Turroncito Tradicional 70 Gr", None, 30, None),
    ("Cantidad", False, "Unidades", "10", "#,##0", 9, None),
    ("Precio Unitario", False, "Precio por unidad en soles", "1.80", "#,##0.00", 14, None),
    ("Total Linea", True, "Cantidad × precio, en soles", "18.00", "#,##0.00", 11, None),
    ("Medio pago", False, "Yape, BCP, BBVA o efectivo", "YAPE", None, 11, ["YAPE", "BCP", "BBVA", "EFECTIVO"]),
    ("Canal", True, "DELIVERY (Lima) o PROVINCIA", "DELIVERY", None, 11, ["DELIVERY", "PROVINCIA"]),
    ("Distrito", False, "Distrito (Lima) o ciudad", "Miraflores", None, 14, None),
    ("Provincia", False, "Provincia (solo provincia)", "Trujillo", None, 13, None),
    ("Departamento", True, "Departamento", "Lima", None, 14, None),
    ("Salió de", False, "Tienda o planta desde donde salió", "Surco", None, 12, None),
    ("Observación", False, "Comentario libre", "", None, 20, None),
]
OXXO = [
    ("FECHA", True, "Día del reporte (dd/mm/aaaa)", "04/10/2026", "dd/mm/yyyy", 12, None),
    ("NOMBRE", True, "Nombre de la tienda OXXO", "Plaza Noviembre", None, 20, None),
    ("DISTRITO", True, "Distrito de la tienda", "SAN ISIDRO", None, 18, None),
    ("CLUSTER", True, "Cluster de OXXO", "A", None, 9, ["A", "B", "C"]),
    ("SUPERGRUPO", False, "Como lo manda OXXO", "PANADERIA INDUSTRIAL", None, 20, None),
    ("CATEGORIA", False, "Como lo manda OXXO", "TURRONES", None, 12, None),
    ("SUBCATEGORIA", False, "Como lo manda OXXO", "TURRONES", None, 13, None),
    ("ESTADO DEL CODIGO", False, "Activo o Inactivo", "Activo", None, 12, ["Activo", "Inactivo"]),
    ("EAN", True, "Código de barras del producto", "733968436658", "@", 15, None),
    ("DESCRIPCION", True, "Producto como lo nombra OXXO", "MiniTurron tradicional calderon 70g", None, 34, None),
    ("Ventas netas(sin IGV)", True, "Venta del día en soles, sin IGV", "9.91", "#,##0.00", 14, None),
    ("Unidades vendidas netas", True, "Unidades vendidas ese día", "3", "#,##0", 14, None),
    ("STOCK", True, "Stock en la tienda al cierre del día", "3", "#,##0", 9, None),
]

PLANTILLAS = {
    "Formato - Reporte interno de tiendas.xlsx": ("Data", TIENDAS,
        "Venta diaria de una o varias tiendas. La hoja de datos DEBE llamarse «Data».",
        "Al cargar se reemplaza, por cada tienda del archivo, desde su primer hasta su último día: subir el mismo archivo dos "
        "veces no duplica la venta."),
    "Formato - Ventas retail.xlsx": ("Ventas", RETAIL,
        "Una fila por despacho a cada cadena (SPSA, OXXO, Tottus, Vendomática, PedidosYa…).",
        "Al cargar se reemplaza, por cada cliente del archivo, desde su primer hasta su último despacho: subir el mismo archivo "
        "dos veces no duplica la venta. Si hay filas idénticas, la vista previa avisa antes de confirmar."),
    "Formato - Ventas virtuales.xlsx": ("2026", VIRTUAL,
        "Una fila por producto de cada comprobante del canal digital (Lima delivery y Provincia). ContaNet es el monto oficial: "
        "este archivo dice de qué canal, zona y producto es cada comprobante.",
        "Al cargar se reemplaza todo el rango de fechas del archivo: subir el mismo archivo dos veces no duplica la venta. Si hay "
        "filas idénticas, la vista previa avisa antes de confirmar."),
    "Formato - Reporte diario OXXO.xlsx": ("Reporte", OXXO,
        "Sell-out que manda OXXO cada día: una fila por tienda y producto con venta, unidades y stock. Normalmente se sube el "
        "archivo de OXXO tal cual; esta plantilla es para armarlo a mano si hiciera falta.",
        "Al cargar se reemplazan los días del archivo (venta por tienda, por producto y stock): subir el mismo día dos veces no "
        "duplica nada. Si una tienda y producto se repite el mismo día, no deja cargar. El «ACUMULADO MENSUAL» no sirve: no trae stock."),
}


def crear(nombre: str, hoja: str, columnas, que_es: str, repetidos: str):
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    for i, (titulo, obligatoria, _, _, formato, ancho, lista) in enumerate(columnas, start=1):
        c = ws.cell(1, i, titulo)
        c.font, c.fill = TITULO, FONDO if obligatoria else FONDO_OPC
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        letra = c.column_letter
        ws.column_dimensions[letra].width = ancho
        for r in range(2, FILAS + 2):
            d = ws.cell(r, i)
            d.font = NORMAL
            if formato:
                d.number_format = formato
        if lista:
            dv = DataValidation(type="list", formula1=f'"{",".join(lista)}"', allow_blank=True,
                                error=f"Elige un valor de la lista: {', '.join(lista)}", errorTitle=titulo, showErrorMessage=True)
            dv.add(f"{letra}2:{letra}{FILAS + 1}")
            ws.add_data_validation(dv)
        elif formato == "dd/mm/yyyy":
            dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2024,1,1)", allow_blank=True,
                                error="Escribe una fecha válida (dd/mm/aaaa).", errorTitle=titulo, showErrorMessage=True)
            dv.add(f"{letra}2:{letra}{FILAS + 1}")
            ws.add_data_validation(dv)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    ins = wb.create_sheet("Instrucciones")
    ins.column_dimensions["A"].width = 24
    ins.column_dimensions["B"].width = 13
    ins.column_dimensions["C"].width = 52
    ins.column_dimensions["D"].width = 30
    ins["A1"] = nombre.removesuffix(".xlsx").removeprefix("Formato - ")
    ins["A1"].font = Font(name="Arial", size=14, bold=True)
    filas = [que_es, f"Llena la hoja «{hoja}» desde la fila 2. Los títulos de la fila 1 no se cambian (el orden de las columnas sí puede cambiar).",
             "Columnas en marrón oscuro: obligatorias. En marrón claro: opcionales.", repetidos]
    for k, t in enumerate(filas, start=2):
        ins.cell(k, 1, t).font = NORMAL
        ins.merge_cells(start_row=k, start_column=1, end_row=k, end_column=4)
        ins.cell(k, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ins.row_dimensions[k].height = 30
    r0 = len(filas) + 3
    for j, t in enumerate(["Columna", "Obligatoria", "Qué poner", "Ejemplo"], start=1):
        c = ins.cell(r0, j, t)
        c.font, c.fill = TITULO, FONDO
    for k, (titulo, obligatoria, que, ejemplo, _, _, lista) in enumerate(columnas, start=r0 + 1):
        detalle = que + (f" ({' / '.join(lista)})" if lista else "")
        for j, v in enumerate([titulo, "Sí" if obligatoria else "No", detalle, ejemplo], start=1):
            c = ins.cell(k, j, v)
            c.font = NEGRITA if j == 1 else NORMAL
            c.border = LINEA
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.active = 0
    DESTINO.mkdir(parents=True, exist_ok=True)
    wb.save(DESTINO / nombre)
    print("Creado:", DESTINO / nombre)


if __name__ == "__main__":
    for nombre, (hoja, columnas, que_es, repetidos) in PLANTILLAS.items():
        crear(nombre, hoja, columnas, que_es, repetidos)
